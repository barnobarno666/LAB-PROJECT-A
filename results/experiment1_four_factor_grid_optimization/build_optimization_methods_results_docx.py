"""Build a short Word paper from the saved four-factor M4 grid search."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SUMMARY = json.loads((HERE / "optimization_summary.json").read_text(encoding="utf-8"))
BEST = SUMMARY["best_grid_case"]
OUTPUT = HERE / "optimization_methods_results.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def set_cell_margin(cell, top=75, start=105, bottom=75, end=105) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, text, end):
        run._r.append(item)


def add_math_run(parent, text: str) -> None:
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    parent.append(run)


def add_subscript(parent, base: str, sub: str) -> None:
    node = OxmlElement("m:sSub")
    main = OxmlElement("m:e")
    add_math_run(main, base)
    lower = OxmlElement("m:sub")
    add_math_run(lower, sub)
    node.extend((main, lower))
    parent.append(node)


def add_conversion_equation(doc: Document) -> None:
    paragraph = doc.add_paragraph(style="Equation")
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    add_subscript(math, "X", "CO,out")
    add_math_run(math, " = 100(1 − ")
    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    denominator = OxmlElement("m:den")
    add_subscript(numerator, "F", "CO,out")
    add_subscript(denominator, "F", "CO,in")
    fraction.extend((numerator, denominator))
    math.append(fraction)
    add_math_run(math, ") %")
    math_para.append(math)
    paragraph._p.append(math_para)


def add_figure(doc: Document, filename: str, caption: str) -> None:
    paragraph = doc.add_paragraph(style="Figure")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(HERE / filename), width=Inches(5.72))
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(caption)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.76)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.37)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16.5)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title.element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    title_rpr = title.element.get_or_add_rPr()
    title_fonts = title_rpr.rFonts
    if title_fonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "csTheme"):
            title_fonts.attrib.pop(qn(f"w:{attr}"), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            title_fonts.set(qn(f"w:{attr}"), "Times New Roman")

    for name in ("Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].font.size = Pt(12.5)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)

    equation_style = styles.add_style("Equation", 1)
    equation_style.base_style = normal
    equation_style.paragraph_format.space_before = Pt(1)
    equation_style.paragraph_format.space_after = Pt(7)
    equation_style.paragraph_format.left_indent = Inches(0.2)

    figure_style = styles.add_style("Figure", 1)
    figure_style.base_style = normal
    figure_style.paragraph_format.space_after = Pt(2)
    figure_style.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption_style = styles["Caption"]
    caption_style.font.name = "Times New Roman"
    caption_style.font.size = Pt(9)
    caption_style.font.bold = False
    caption_style.font.color.rgb = RGBColor(0, 0, 0)
    caption_style.paragraph_format.line_spacing = 1.05
    caption_style.paragraph_format.space_before = Pt(1)
    caption_style.paragraph_format.space_after = Pt(12)
    caption_style.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = normal
    footer.add_run("Page ")
    add_field(footer, "PAGE")
    for run in footer.runs:
        run.font.size = Pt(8)

    doc.add_paragraph("Four Factor Optimization of CO Methanation in a Fixed Bed Reactor", style="Title")
    doc.add_paragraph(
        "A 20,412-case grid search of the Full M4 reactor model gave a best sampled setting of "
        "370 °C, 15 bar absolute, H₂/CO = 5.0, and 140.4 kg catalyst·s/mol CO. The predicted "
        "outlet CO conversion was 99.999993%. This is a model result for the stated grid."
    )

    doc.add_heading("Methodology", level=1)
    doc.add_paragraph(
        "The search used the Full M4 fixed bed model on the Experiment 1 basis. Catalyst mass was "
        "3.12 g, and the inlet N₂/CO molar ratio was fixed at 2.496875. The inherited reactor geometry "
        "and an assumed particle density of 1250 kg m⁻³ give a bed voidage of 0.98384. Each case was "
        "integrated to the bed outlet under an isothermal assumption with Ergun pressure drop "
        "and the BDF solver."
    )
    doc.add_paragraph(
        "The grid varied inlet H₂/CO ratio, temperature, absolute pressure, and catalyst space time "
        "on a CO feed basis. Its 12 × 9 × 9 × 21 combinations all completed. At fixed catalyst mass, "
        "space time set the CO inlet flow; the H₂ and N₂ flows followed from their respective inlet "
        "ratios. Total inlet flow therefore changed with space time and H₂/CO ratio."
    )
    doc.add_paragraph(
        "The objective was the completed bed outlet CO conversion, defined as"
    )
    add_conversion_equation(doc)
    doc.add_paragraph(
        "The highest conversion was selected, with no hydrogen, pressure, or throughput cost in the "
        "objective. Each figure varies one factor across its saved values while the other three "
        "remain at the selected point. The pressure grid covers 1–15 bar absolute. Kinetic parameters "
        "were fitted over 5–15 bar, so the 1–4 bar predictions are extrapolations."
    )

    doc.add_heading("Results and discussion", level=1)
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.keep_with_next = True
    caption.add_run("Table 1. Search ranges and best sampled model result.")

    rows = [
        ("Inlet H₂/CO ratio", "1.0–5.0 (12 levels)", f"{BEST['h2_co_molar_ratio']:.1f}"),
        ("Inlet temperature", "250–400 °C (9 levels)", f"{BEST['inlet_temperature_c']:.0f} °C"),
        ("Inlet pressure", "1–15 bar abs (9 levels)", f"{BEST['inlet_pressure_bar_abs']:.0f} bar abs"),
        ("Catalyst space time", "8.775–140.4 (21 levels)", "140.4 kg catalyst·s/mol CO"),
        ("CO / H₂ / N₂ inlet flow", "Derived from inputs", "0.080 / 0.400 / 0.19975 mol h⁻¹"),
        ("Outlet CO conversion", "Objective", f"{BEST['outlet_co_conversion_pct']:.6f}%"),
        ("Outlet CH₄ yield", "Predicted output", f"{BEST['outlet_ch4_yield_pct']:.6f}%"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.22), Inches(2.05), Inches(2.42)]
    for cell, name, width in zip(table.rows[0].cells, ("Quantity", "Grid or role", "Best sampled value"), widths):
        cell.width = width
        cell.text = name
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = width
            cell.text = value
            set_cell_shading(cell, "F4F5F6" if row_index % 2 else "FFFFFF")
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margin(cell)
            if row_index == 0:
                set_cell_shading(cell, "263746")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.07
                if col_index > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.8)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)

    analysis = doc.add_paragraph(
        "The selected temperature was an interior grid value; H₂/CO ratio, pressure, and space "
        "time reached their upper sampled bounds. At the selected values of the other inputs, raising "
        "H₂/CO from 1.0 to 1.4 increased conversion from 89.27% to 99.14% (Figure 1). Subsequent "
        "gains were smaller. Temperature peaked at the sampled 370 °C point, with a slight decline "
        "at 385 and 400 °C (Figure 2)."
    )
    analysis.paragraph_format.space_before = Pt(5)
    doc.add_paragraph(
        "The pressure and space time slices also approach a plateau (Figures 3 and 4). Over the "
        "fitted pressure interval, conversion changed from 99.999112% at 5 bar to 99.999993% at "
        "15 bar. Increasing space time from 8.775 to 140.4 kg catalyst·s/mol CO raised conversion "
        "from 99.5533% to 99.999993%. At the selected point, CO feed was 0.080 mol h⁻¹. At the "
        "nominal Experiment 1 inputs, the model predicts 96.94% conversion against a source-reported "
        "40.63%. This unresolved difference and the assumed bed geometry limit interpretation of the "
        "optimized model result."
    )

    doc.add_page_break()
    add_figure(doc, "sensitivity_h2_co_ratio.png", "Figure 1. Outlet CO conversion across the sampled H₂/CO ratios.")
    add_figure(doc, "sensitivity_inlet_temperature.png", "Figure 2. Outlet CO conversion across the sampled inlet temperatures.")

    doc.add_page_break()
    add_figure(doc, "sensitivity_inlet_pressure.png", "Figure 3. Outlet CO conversion across the sampled inlet pressures.")
    add_figure(doc, "sensitivity_catalyst_space_time.png", "Figure 4. Outlet CO conversion across the sampled catalyst space times.")

    doc.core_properties.title = "Four Factor Optimization of CO Methanation in a Fixed Bed Reactor"
    doc.core_properties.subject = "Methodology and results from the saved Full M4 grid search"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
