from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "output" / "paper"
FIGURES = PAPER / "figures"
OUTPUT = PAPER / "fixed_bed_co_methanation_paper.docx"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def set_run_font(run, name: str, size: float, *, bold: bool = False,
                 italic: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=95, start=105, bottom=95, end=105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D7DC", size="5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def set_document_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.widow_control = True

    for style_name, size, before, after in (
        ("Heading 1", 15, 12, 5),
        ("Heading 2", 11.5, 9, 3.5),
    ):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(8.7)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.font.bold = False
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.line_spacing = 1.05
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_body(doc: Document, text: str, *, first_indent: bool = True) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Inches(0.18) if first_indent else Inches(0)
    p.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.style = doc.styles["List Bullet"]
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(2.8)
    p.add_run(text)


def add_equation(doc: Document, text: str, number: int | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.48), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    run = p.add_run(text)
    set_run_font(run, "Cambria Math", 9.2)
    if number is not None:
        p.add_run("\t")
        run = p.add_run(f"({number})")
        set_run_font(run, "Times New Roman", 9)


def add_caption(doc: Document, label: str, text: str, *, table=False) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{label}. ")
    r.bold = True
    p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "FFFFFF")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.02
        r = p.add_run(header)
        set_run_font(r, "Times New Roman", 8.6, bold=True, color="000000")
    repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            r = p.add_run(value)
            set_run_font(r, "Times New Roman", 8.15, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, filename: str, number: int, caption: str,
               width: float = 6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_together = True
    r = cap.add_run(f"Figure {number}. ")
    r.bold = True
    cap.add_run(caption)


def add_reference(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"[{number}] ")
    r.bold = True
    set_run_font(p.add_run(text), "Times New Roman", 8.6)


def make_paper() -> None:
    PAPER.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.32)
    set_document_styles(doc)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("CO Methanation in a Fixed-Bed Reactor"),
                 "Times New Roman", 8, color="333333")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    set_run_font(footer.add_run(""), "Times New Roman", 8, color="555555")
    add_field(footer, "PAGE")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run("One-Dimensional Modeling of CO Methanation in a Fixed-Bed Reactor"),
                 "Times New Roman", 20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run("Barno and Ashiq"), "Times New Roman", 10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run("Abstract"), "Times New Roman", 10.5, bold=True)
    abstract = (
        "A one-dimensional fixed-bed model was used to evaluate CO methanation with the Full M4 "
        "kinetics of Celoria et al. Species, energy, and Ergun pressure balances were integrated "
        "in Python with SciPy solve_ivp using Radau or BDF. At 350 °C, 1 bar absolute, and 3.12 g "
        "catalyst, Full M4 predicts 78.72% CO conversion without parameter adjustment; the "
        "reported Experiment 1 value is 40.63%. This comparison is affected by conflicting "
        "pressure values in the source report and lies below the model's 5–15 bar fitted "
        "pressure range. A separate assumed 5 bar bed predicts 95.81% outlet conversion and a "
        "peak gas temperature of 819.7 K (546.5 °C), above the fitted temperature range of "
        "250–400 °C. Sweeps of inlet temperature, pressure, H₂/CO ratio, catalyst space time, "
        "and N₂ flow characterize the model's response under specified conditions."
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.06
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(p.add_run(abstract), "Times New Roman", 9.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Keywords  ")
    set_run_font(r, "Times New Roman", 9.2, bold=True)
    set_run_font(p.add_run("CO methanation; fixed-bed reactor; Full M4 kinetics; stiff ODE integration; sensitivity analysis"),
                 "Times New Roman", 9.2)

    doc.add_heading("Introduction", level=1)
    add_body(doc,
        "CO methanation is strongly exothermic, so local rates change with gas composition and "
        "temperature along a packed bed. A reactor calculation must also account for heat "
        "exchange and pressure loss."
    )
    add_body(doc,
        "Here, the Full M4 kinetics of Celoria et al. [1] are applied to a one-dimensional "
        "fixed-bed reactor. Its unadjusted prediction at the selected Experiment 1 basis is "
        "compared with two other literature rate laws [2, 3] and the reported conversion [5]. "
        "A separate assumed 5 bar case and operating sweeps are used to examine axial behavior "
        "and input sensitivity."
    )

    doc.add_heading("Methodology", level=1)
    doc.add_heading("Reactor model and reaction network", level=2)
    add_body(doc,
        "The bed is represented as a steady, one-dimensional plug-flow reactor. The Full M4 "
        "network contains three reversible reactions, written below in the forward directions "
        "used to define the stoichiometric matrix. The published reverse water-gas-shift rate "
        "is sign-reversed for the forward direction used here."
    )
    for reaction in (
        "CO₂ + 4 H₂ ⇌ CH₄ + 2 H₂O",
        "CO + 3 H₂ ⇌ CH₄ + H₂O",
        "CO + H₂O ⇌ CO₂ + H₂",
    ):
        add_bullet(doc, reaction)
    add_body(doc,
        "Full M4 local rates are evaluated from temperature and gas partial pressures using "
        "the published Langmuir–Hinshelwood–Hougen–Watson form and central parameter values "
        "[1]. For comparison, the Kopyscinski implementation uses the corrected square-root "
        "CO adsorption term from its published corrigendum [2]. The Quindimil implementation "
        "was fitted to CO₂/H₂ feeds [3]; applying it to the CO/H₂/N₂ comparison is outside that "
        "feed scope."
    )

    doc.add_heading("Material, energy, and momentum balances", level=2)
    add_body(doc,
        "Species molar flows Fᵢ are integrated with catalyst mass W as the independent "
        "coordinate. The activity factor a multiplies each net reaction rate rⱼ. Local partial "
        "pressures follow the ideal-gas mixture relation pᵢ = yᵢP, and concentrations are "
        "calculated from cᵢ = pᵢ/(RT)."
    )
    add_equation(doc, "dFᵢ/dW = a ∑ⱼ νᵢⱼ rⱼ", 1)
    add_body(doc,
        "For non-isothermal cases, the gas energy balance includes reaction heat and, when "
        "enabled, heat transfer to a wall held at T_w. The model uses the species heat-capacity "
        "flow in the denominator and a cylindrical heat-transfer area per bed volume of 4/D_t."
    )
    add_equation(doc,
        "dT/dW = [−a ∑ⱼ ΔHⱼ rⱼ − (4U/(Dₜρ_cat))(T − T_w)] / [∑ᵢ Fᵢ Cₚ,ᵢ]", 2)
    add_body(doc,
        "Pressure loss is computed from the viscous and inertial terms of the Ergun equation. "
        "The axial pressure gradient is converted to the catalyst-mass coordinate using the "
        "bed cross-sectional area A_c and catalyst loading per bed volume ρ_cat."
    )
    add_equation(doc,
        "dP/dW = −[150μ(1−ε)²u/(dₚ²ε³) + 1.75ρ_g(1−ε)u²/(dₚε³)]/(ρ_cat A_c)", 3)
    add_equation(doc,
        "X_CO = 100(F_CO,0 − F_CO,out)/F_CO,0;   τ_cat = W_cat/F_CO,0", 4)

    doc.add_heading("Numerical solution and simulation cases", level=2)
    add_body(doc,
        "The Python implementation integrates the balance equations with "
        "scipy.integrate.solve_ivp from SciPy 1.18.1 [4], using the implicit Radau or BDF "
        "method as listed in Table 1. The principal profile runs use relative tolerance 10⁻⁶, "
        "absolute "
        "flow tolerance 10⁻¹⁰ mol s⁻¹, temperature tolerance 10⁻⁵ K, and pressure tolerance "
        "0.1 Pa. The independent literature-model profiles use Radau with relative tolerance "
        "10⁻⁸ and absolute flow tolerance 10⁻¹⁴ mol s⁻¹. Dense solver output is sampled at 401 "
        "positions for reactor profiles and 101 catalyst-mass positions for the three-model "
        "comparison. Sweep results are reported at the bed outlet."
    )
    add_body(doc,
        "A dry CO/H₂ inlet has zero initial H₂O, while the Full M4 CO₂ rate contains a water "
        "pressure denominator. The implementation uses a derived near-inlet series start for "
        "the stiff integration rather than adding an artificial steam feed."
    )

    add_caption(doc, "Table 1", "Simulation cases represented in the figures.", table=True)
    add_table(doc,
        ["Case", "Conditions and varied inputs", "Integrator and coverage"],
        [
            ["Three-model comparison", "350 °C; 1 bar abs.; 3.12 g; fixed Experiment 1 inlet flows; isothermal", "Radau; 101 profile points per model"],
            ["Assumed Full M4 profile", "350 °C; 5 bar; CO:H₂:N₂ = 1:4:25; heat exchange; Ergun drop", "Radau; 401 axial output points"],
            ["Temperature–pressure sweep", "22 × 22 grid; inlet temperature 250–400 °C; pressure 1–15 bar", "Radau; 484 cases completed"],
            ["Composition sensitivities", "H₂/CO–space time, feed-flow and isolated N₂ sweeps at 1 bar", "BDF; 462, 819 and 42 cases completed"],
        ],
        [1.42, 3.54, 1.50]
    )
    add_body(doc,
        "The Full M4 kinetic fit covers 250–400 °C and 5–15 bar. The 1 bar Experiment 1 "
        "comparison and sensitivity runs are pressure extrapolations. The assumed reactor case "
        "uses a 0.0254 m "
        "tube, 0.3048 m bed, 0.4 void fraction, 3 mm particles, and 750 kg catalyst m⁻³ bed; "
        "these are specified modeling inputs, not dimensions or packing measurements confirmed "
        "by Experiment 1."
    )

    doc.add_heading("Results and Discussion", level=1)
    doc.add_heading("Comparison with Experiment 1", level=2)
    add_body(doc,
        "The common comparison basis is 350 °C, 1.00 bar absolute, 3.12 g catalyst, and inlet "
        "flows of 0.320 mol h⁻¹ CO, 0.959 mol h⁻¹ H₂, and 0.799 mol h⁻¹ N₂. The source reports "
        "40.63% CO conversion [5]. No rate or activity parameter was fitted to that value."
    )
    add_caption(doc, "Table 2", "Independent model predictions at the common Experiment 1 basis.", table=True)
    add_table(doc,
        ["Model", "Predicted conversion", "Difference", "Scope at 1 bar"],
        [
            ["Full M4", "78.723%", "+38.09 percentage points", "Below the 5–15 bar fit range"],
            ["Kopyscinski", "99.985%", "+59.35 percentage points", "Within the published 1–2 bar range"],
            ["Quindimil", "3.369%", "−37.26 percentage points", "Below the 2–6 bar range; CO/H₂ feed outside its CO₂/H₂ fit scope"],
        ],
        [1.08, 1.18, 1.43, 2.77]
    )
    add_body(doc,
        "Table 2 and Figure 1 show that Full M4 predicts 78.72% conversion at the selected "
        "1 bar basis, 38.09 percentage points above the reported 40.63%. The source describes "
        "atmospheric operation in its abstract and discussion but lists 2 bar in Appendix A [5]. "
        "The 1 bar calculation is below the Full M4 fitted pressure interval."
    )
    add_body(doc,
        "Kopyscinski predicts nearly complete conversion, but it describes a different catalyst "
        "and reactor. Quindimil's rate law was fitted to CO₂/H₂ feeds and is "
        "evaluated below its pressure range here. Its implemented expression is also reduced "
        "because the published parameter table lacks an independent K_H₂ required by one "
        "denominator term. These results assess transfer of literature kinetics across different "
        "conditions; they do not rank models for a common catalyst."
    )
    add_figure(doc, "figure_01_model_comparison.png", 1,
        "CO conversion versus catalyst mass for three literature models evaluated without refitting. The dashed line and circular point mark the reported Experiment 1 conversion at 3.12 g.")

    doc.add_heading("Assumed Full M4 reactor profile", level=2)
    add_body(doc,
        "Figure 2 shows the assumed 5 bar Full M4 bed described in Table 1. The total inlet flow "
        "is 0.05 mol s⁻¹ with CO:H₂:N₂ = 1:4:25; the wall is held at 350 °C with U = 10 W m⁻² K⁻¹. "
        "At 0.11583 kg catalyst, the model predicts 95.81% outlet CO conversion and 87.33% "
        "methane yield. Gas temperature peaks at 819.66 K (546.5 °C) about 3.5 cm from the "
        "inlet, while pressure falls to 4.9405 bar at the outlet."
    )
    add_figure(doc, "figure_02_assumed_reactor_profiles.png", 2,
        "Axial CO conversion, gas temperature and pressure for the assumed Full M4 case. The temperature rise is concentrated near the inlet; the small pressure change is plotted on a separate right-hand axis. A circular point marks the inlet condition in each panel.")
    add_body(doc,
        "In Figure 3, CO depletion coincides with CH₄ and H₂O formation near the inlet. CO₂ "
        "remains a minor intermediate; N₂ is included in the material balance but omitted from "
        "the plot scale. The calculated hot spot exceeds the 400 °C upper kinetic-fit temperature, "
        "so rates in part of the bed require temperature extrapolation."
    )
    add_figure(doc, "figure_03_species_profiles.png", 3,
        "Reactive-species gas concentrations along the assumed Full M4 bed. The logarithmic ordinate omits initial zero concentrations; nitrogen is omitted from the panel for scale clarity and remains in the simulated material balance.")

    doc.add_heading("Temperature and pressure sensitivity", level=2)
    tp_rows = read_csv(ROOT / "results" / "temperature_pressure_full_m4_sweep" /
                       "temperature_pressure_sweep.csv")
    fit_rows = [r for r in tp_rows if r["status"] == "completed" and
                float(r["inlet_pressure_bar_abs"]) >= 5.0]
    fit_values = [float(r["outlet_co_conversion_pct"]) for r in fit_rows]
    add_body(doc,
        f"Figure 4 shows the 5–15 bar subset of a 22 × 22 Full M4 sweep over 250–400 °C "
        f"and 1–15 bar. All 484 runs completed. Over the plotted pressure interval, outlet "
        f"conversion ranges from {min(fit_values):.2f}% to {max(fit_values):.2f}%; the nominal "
        f"350 °C, 5 bar point gives 95.81%. Cases below 5 bar are omitted because they lie "
        f"outside the fitted pressure range."
    )
    add_figure(doc, "figure_04_temperature_pressure_map.png", 4,
        "Full M4 outlet CO conversion over the 5–15 bar pressure-fit interval and 250–400 °C inlet-temperature range. The circular marker identifies the 350 °C, 5 bar nominal case. Other assumed bed inputs are held fixed.")
    add_body(doc,
        "Inlet conditions within the fitted interval do not guarantee that the entire bed stays "
        "there. Across all 484 runs, including those below 5 bar, 469 reach a sampled peak bed "
        "temperature above 400 °C; peak temperatures range from 288.4 to 603.1 °C. Many plotted "
        "trajectories therefore involve local temperature extrapolation."
    )

    doc.add_heading("H₂/CO ratio and catalyst space time", level=2)
    ratio_meta = read_json(ROOT / "results" / "experiment1_ratio_contact_time_sweep_1bar" /
                           "sweep_metadata.json")
    add_body(doc,
        "Figure 5 summarizes a 462-case isothermal Full M4 sweep of H₂/CO from 1 to 5 and "
        "catalyst space time "
        "from 8.775 to 140.4 kg_cat s mol_CO⁻¹ at 350 °C and 1 bar. Catalyst mass was held at "
        "3.12 g; CO and N₂ flows were scaled together, H₂ was reset from the selected ratio, and "
        "N₂/CO was held at the Experiment 1 value. Outlet conversion ranges from "
        f"{ratio_meta['conversion_range_pct'][0]:.2f}–{ratio_meta['conversion_range_pct'][1]:.2f}%. "
        "At the reference ratio of 2.9969 and space time of 35.10 kg_cat s mol_CO⁻¹, the model "
        "returns 78.72%. Increasing H₂/CO or space time generally raises conversion over the "
        "sampled region, with diminishing change toward the high-conversion boundary."
    )
    add_figure(doc, "figure_05_ratio_space_time_map.png", 5,
        "Full M4 outlet conversion across the H₂/CO and catalyst-space-time grid. The circular point marks the 1 bar Experiment 1 reference calculation. All cells are isothermal predictions without refitting and lie below the 5–15 bar Full M4 pressure-fit interval.")

    doc.add_heading("Absolute feed-flow sensitivity", level=2)
    add_body(doc,
        "Figure 6 varies absolute inlet flows at 350 °C, 1 bar, and 3.12 g catalyst. Panel (a) "
        "varies H₂ and CO with N₂ fixed at 0.799 mol h⁻¹; panel (b) varies H₂ and N₂ with CO "
        "fixed at 0.320 mol h⁻¹. The two grids contain 819 completed cases, with predicted "
        "conversion from 15.91% to 92.30%. The pressure is below the Full M4 fitted interval."
    )
    add_figure(doc, "figure_06_feed_flow_sensitivities.png", 6,
        "Full M4 feed-flow sensitivity at 1 bar. Panel (a) varies CO and H₂ at fixed N₂; panel (b) varies N₂ and H₂ at fixed CO. Circular points mark the Experiment 1 inlet flows; both panels share a color scale.")
    add_body(doc,
        "At fixed N₂, conversion rises with H₂ supply relative to CO. In panel (b), added N₂ "
        "lowers conversion at fixed reactive flows; dilution and total inlet flow both change."
    )

    doc.add_heading("Nitrogen-flow sensitivity at fixed reactive flows", level=2)
    add_body(doc,
        "Figure 7 shows a 42-case N₂ sweep with CO held at 0.320 mol h⁻¹ and H₂ at 0.959 mol h⁻¹. "
        "As N₂ varies from 0.400 to 1.200 mol h⁻¹, predicted conversion falls from 88.14% to 70.34%; "
        "the reference N₂ flow of 0.799 mol h⁻¹ gives 78.72%. Since reactive flows are fixed, "
        "the curve captures the combined effect of dilution and increased total flow at 1 bar, "
        "below the fitted pressure range."
    )
    add_figure(doc, "figure_07_nitrogen_sensitivity.png", 7,
        "Full M4 conversion as N₂ flow changes with CO and H₂ held at their Experiment 1 values. The circular marker identifies the reference N₂ feed. The model is evaluated isothermally at 350 °C and 1 bar.")

    doc.add_heading("Conclusions", level=1)
    add_body(doc,
        "At the selected 1 bar Experiment 1 conditions, Full M4 predicts 78.72% CO conversion, "
        "compared with the reported 40.63%. Because 1 bar lies below the fitted pressure range "
        "and the source gives conflicting pressure values, this difference does not establish "
        "predictive accuracy. In the separate assumed 5 bar bed, conversion reaches 95.81%, but "
        "the calculated 546.5 °C hot spot exceeds the published temperature-fit range. Over the "
        "specified sensitivity cases, conversion rises with H₂/CO and catalyst space time and "
        "falls as N₂ flow increases at fixed CO and H₂. An experimental comparison requires a "
        "resolved pressure basis, measured reactor geometry and packing, and kinetics for the "
        "catalyst and local bed conditions."
    )
    doc.paragraphs[-1].paragraph_format.keep_together = True

    doc.add_heading("References", level=1)
    add_reference(doc, 1,
        "F. Celoria, F. Salomone, A. Tauro, et al., “Kinetic study and deactivation phenomena for the methanation of CO₂ and CO mixed syngas on a Ni/Al₂O₃ catalyst,” Chemical Engineering Journal, vol. 512, article 162113, 2025. doi: 10.1016/j.cej.2025.162113.")
    add_reference(doc, 2,
        "J. Kopyscinski, T. J. Schildhauer, F. Vogel, S. M. A. Biollaz, and A. Wokaun, “Applying spatially resolved concentration and temperature measurements in a catalytic plate reactor for the kinetic study of CO methanation,” Journal of Catalysis, vol. 271, no. 2, pp. 262–279, 2010. doi: 10.1016/j.jcat.2010.02.008. See also the corrigendum, Journal of Catalysis, vol. 273, no. 1, p. 82, 2010. doi: 10.1016/j.jcat.2010.05.001.")
    add_reference(doc, 3,
        "A. Quindimil, J. A. Onrubia-Calvo, A. Davó-Quiñonero, et al., “Intrinsic kinetics of CO₂ methanation on low-loaded Ni/Al₂O₃ catalyst: Mechanism, model discrimination and parameter estimation,” Journal of CO₂ Utilization, vol. 57, article 101888, 2022. doi: 10.1016/j.jcou.2022.101888.")
    add_reference(doc, 4,
        "P. Virtanen, R. Gommers, T. E. Oliphant, et al., “SciPy 1.0: Fundamental algorithms for scientific computing in Python,” Nature Methods, vol. 17, pp. 261–272, 2020. doi: 10.1038/s41592-019-0686-2.")
    add_reference(doc, 5,
        "Experiment 1 source report and data supplied with the project, unpublished project material.")

    props = doc.core_properties
    props.title = "One-Dimensional Modeling of CO Methanation in a Fixed-Bed Reactor"
    props.subject = "Reactor methodology, profiles, and operating sensitivities"
    props.author = "Barno and Ashiq"
    props.keywords = "CO methanation, fixed bed, Full M4, SciPy, Radau, BDF"
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    make_paper()
